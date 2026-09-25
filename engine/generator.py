"""
Core engine for the SBA Result Card generator.

Pipeline for each student row found in the uploaded Log Sheet:
  1. Stamp that row's data into the bundled master template's
     'Result Card'!A41:Q41 staging row (exactly what the original
     VBA macro did).
  2. Hide the 'Log Sheet' tab so only the Result Card prints/exports.
  3. Convert the workbook to PDF with LibreOffice headless, which
     recalculates every formula (grades, percentages, grand total)
     using the template's own, unmodified formulas.

No formulas, formatting, layout, or wording in the template are
touched -- only the 17 input cells in the staging row change.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = BASE_DIR / "assets" / "master_template.xlsx"

# Column order the Result Card staging row (A41:Q41) expects.
STAGING_COLUMNS = [
    "school", "class", "roll", "bform", "name", "father", "dob", "section",
    "english", "urdu", "maths", "gk", "social", "islamiat", "quran",
    "computer", "other",
]

SUBJECT_KEYS = [
    "english", "urdu", "maths", "gk", "social", "islamiat", "quran",
    "computer", "other",
]

# Accepted header text (normalized: lowercased, collapsed whitespace, no
# trailing colon) for each field, so the uploaded file doesn't have to
# match the original template's headers character-for-character.
HEADER_SYNONYMS: dict[str, list[str]] = {
    "school": ["school name"],
    "class": ["class"],
    "roll": ["roll no", "roll no.", "rollno", "roll number"],
    "bform": ["b.form no", "b.form no.", "bform no", "b form no", "b-form no", "b.form number"],
    "name": ["student name"],
    "father": ["father name"],
    "dob": ["date of birth", "d.o.b", "d.o.b.", "dob"],
    "section": ["section"],
    "english": ["english"],
    "urdu": ["urdu"],
    "maths": ["mathematics", "maths", "math"],
    "gk": [
        "general knowledge/general science",
        "general knowledge / general science",
        "gk/gs", "gk / gs",
    ],
    "social": [
        "social studies/history/geoghraghy",
        "social studies/history/geography",
        "social studies / history / geography",
    ],
    "islamiat": ["islamiat/ethics", "islamiat / ethics", "islamiat"],
    "quran": ["holy quran"],
    "computer": ["computer education"],
    "other": ["other subject"],
}

REQUIRED_KEYS = ["name", "class", "section"]


class ParseError(Exception):
    """Raised for problems with the uploaded file that the user can fix."""


def _norm(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.rstrip(":")
    text = re.sub(r"\s+", " ", text)
    return text


def _build_header_lookup() -> dict[str, str]:
    lookup = {}
    for key, variants in HEADER_SYNONYMS.items():
        for variant in variants:
            lookup[_norm(variant)] = key
    return lookup


HEADER_LOOKUP = _build_header_lookup()


def _find_log_sheet(wb: openpyxl.Workbook):
    """Prefer a sheet literally named like 'Log Sheet'; otherwise pick
    whichever sheet's header row matches the most known field names."""
    for name in wb.sheetnames:
        if "log" in name.lower():
            return wb[name]

    best_ws, best_score = None, -1
    for name in wb.sheetnames:
        ws = wb[name]
        header_row = next(ws.iter_rows(min_row=1, max_row=1), [])
        score = sum(1 for cell in header_row if _norm(cell.value) in HEADER_LOOKUP)
        if score > best_score:
            best_ws, best_score = ws, score
    if best_ws is None:
        raise ParseError("Could not find any sheet with recognizable columns in the uploaded file.")
    return best_ws


def _map_headers(ws) -> dict[str, int]:
    """Map field key -> 1-based column index, using the first matching
    column for each key (handles the template's duplicated 'Section'
    column -- the first one, which also appears on the printed card,
    wins; a later duplicate is ignored)."""
    col_map: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col).value
        key = HEADER_LOOKUP.get(_norm(header_val))
        if key and key not in col_map:
            col_map[key] = col

    missing = [k for k in REQUIRED_KEYS if k not in col_map]
    if missing:
        raise ParseError(
            "Missing required column(s) in the uploaded sheet: "
            + ", ".join(missing)
            + ". The uploaded file should use the same column headers as the "
              "original Log Sheet template."
        )
    return col_map


def _num(value: Any) -> int | float:
    if value is None or value == "":
        return 0
    try:
        f = float(value)
    except (TypeError, ValueError):
        return 0
    return int(f) if f.is_integer() else f


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def extract_students(upload_path: Path) -> list[dict[str, Any]]:
    """Read every student row from the uploaded workbook's Log Sheet."""
    wb = openpyxl.load_workbook(upload_path, data_only=True)
    ws = _find_log_sheet(wb)
    col_map = _map_headers(ws)

    students = []
    for row in range(2, ws.max_row + 1):
        name_val = ws.cell(row=row, column=col_map["name"]).value
        if not _text(name_val):
            continue  # skip the totals row, blank rows, etc.

        def get(key):
            col = col_map.get(key)
            return ws.cell(row=row, column=col).value if col else None

        student = {
            "school": _text(get("school")),
            "class": _num(get("class")) or "Unknown_Class",
            "roll": _num(get("roll")),
            "bform": _text(get("bform")),
            "name": _text(name_val),
            "father": _text(get("father")),
            "dob": get("dob"),
            "section": _text(get("section")) or "Unknown_Section",
        }
        for key in SUBJECT_KEYS:
            student[key] = _num(get(key))
        students.append(student)

    if not students:
        raise ParseError(
            "No student rows were found. Make sure data starts on the row right "
            "after the header row, with the Student Name column filled in."
        )
    return students


def _sanitize(value: Any) -> str:
    text = str(value).strip()
    text = re.sub(r'[\\/:*?"<>|]', "-", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or "Unknown"


def _student_filename(student: dict[str, Any]) -> str:
    return f"{_sanitize(student['name'])}_{_sanitize(student['class'])}_{_sanitize(student['roll'])}.pdf"


def _render_pdf(student: dict[str, Any], out_pdf_path: Path, workdir: Path) -> None:
    wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False, keep_vba=False)
    rc = wb["Result Card"]
    ls = wb["Log Sheet"]

    for col_letter, key in zip("ABCDEFGHIJKLMNOPQ", STAGING_COLUMNS):
        rc[f"{col_letter}41"] = student[key]

    # Hide the Log Sheet so the PDF export only contains the Result Card.
    ls.sheet_state = "hidden"

    stamped_path = workdir / f"{uuid.uuid4().hex}.xlsx"
    wb.save(stamped_path)

    lo_profile = workdir / f"lo_profile_{uuid.uuid4().hex}"
    try:
        subprocess.run(
            [
                "soffice", "--headless", "--norestore", "--invisible",
                f"-env:UserInstallation=file://{lo_profile}",
                "--convert-to", "pdf",
                "--outdir", str(workdir),
                str(stamped_path),
            ],
            check=True, timeout=90, capture_output=True,
        )
    except subprocess.CalledProcessError as e:
        raise RuntimeError(
            f"LibreOffice failed to render a PDF for {student['name']!r}: "
            f"{e.stderr.decode(errors='ignore')[:500]}"
        ) from e
    finally:
        shutil.rmtree(lo_profile, ignore_errors=True)

    produced = stamped_path.with_suffix(".pdf")
    if not produced.exists():
        raise RuntimeError(f"PDF was not produced for {student['name']!r}.")
    shutil.move(str(produced), str(out_pdf_path))
    stamped_path.unlink(missing_ok=True)


def generate(upload_path: Path, workdir: Path) -> tuple[str, Path, str]:
    """Process the uploaded file and produce either a single PDF or a ZIP.

    Returns (kind, path, download_filename) where kind is "pdf" or "zip".
    Both the PDF/ZIP and any scratch files are written inside `workdir`,
    which the caller is responsible for cleaning up.
    """
    students = extract_students(upload_path)

    if len(students) == 1:
        student = students[0]
        filename = _student_filename(student)
        out_path = workdir / filename
        _render_pdf(student, out_path, workdir)
        return "pdf", out_path, filename

    zip_path = workdir / "Result_Cards.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for student in students:
            filename = _student_filename(student)
            pdf_path = workdir / filename
            _render_pdf(student, pdf_path, workdir)
            arcname = f"{_sanitize(student['class'])}/{_sanitize(student['section'])}/{filename}"
            zf.write(pdf_path, arcname)
            pdf_path.unlink()
    return "zip", zip_path, "Result_Cards.zip"
